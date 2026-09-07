"""Хранение сырого журнала фоновой сессии и его экспорт в документы.

Сырой журнал пишется построчно и сразу сбрасывается на диск: если приложение
или устройство упадёт, накопленное останется на месте. Файл делится на части
по размеру и числу строк, чтобы его можно было открыть обычным редактором.
Экспорт в txt/docx/xlsx выполняется при остановке сессии и имеет собственные
пороги: одно и то же число строк ведёт себя в Word и Excel по-разному.
"""
from __future__ import annotations

import re
import shutil
from datetime import datetime
from pathlib import Path

from ven4control.paths import LOG_DIR


# Порог сырого файла: 20 МБ комфортно открываются текстовым редактором.
RAW_MAX_BYTES = 20 * 1024 * 1024
RAW_MAX_LINES = 200_000

# DOCX раздувается на XML-разметке в разы и заметно тормозит в Word.
DOCX_MAX_LINES = 5_000

# Лист Excel формально держит больше миллиона строк, но тормозит задолго
# до лимита, поэтому ограничение взято с запасом.
XLSX_MAX_ROWS = 50_000

EXPORT_FORMATS: tuple[str, ...] = ("txt", "docx", "xlsx")


def safe_name(name: str) -> str:
    """Имя устройства, пригодное для имени папки и файла.

    Кириллица разрешена явно: большинство устройств в проекте названы
    по-русски ("Роутер Ишим" и т.п.), а без диапазона а-яА-ЯёЁ такое имя
    целиком состоит из запрещённых символов и схлопывается в generic
    "device" — два разных кириллических устройства тогда писали бы сессии
    в одну и ту же папку, затирая журналы друг друга (найдено живым тестом
    с реальным вторым устройством, не только юнит-тестом на одно имя).
    """
    return re.sub(r"[^A-Za-zА-Яа-яЁё0-9_.-]+", "_", name).strip("_") or "device"


def session_stamp(moment: datetime | None = None) -> str:
    return (moment or datetime.now()).strftime("%Y%m%d-%H%M%S")


def _device_folder(device_name: str, device_id: int | None) -> str:
    """Имя папки устройства: имя + id, если он есть.

    id — единственное, что гарантированно уникально между устройствами:
    имя пользователь может задать любое, в т.ч. одинаковое двум разным
    устройствам (или оба схлопнутся в один и тот же safe_name, как было
    с "Роутер" и "Роутер!!!" до появления id в суффиксе). Без id вторая
    сессия просто дописывалась бы в папку первой.
    """
    base = safe_name(device_name)
    return f"{base}-{device_id}" if device_id is not None else base


def session_directory(
    device_name: str,
    moment: datetime | None = None,
    root: Path = LOG_DIR,
    device_id: int | None = None,
) -> Path:
    """Папка одной сессии: <root>/<устройство>[-id]/<дата-время>."""
    return Path(root) / _device_folder(device_name, device_id) / session_stamp(moment)


def session_base_name(
    device_name: str,
    moment: datetime | None = None,
    device_id: int | None = None,
) -> str:
    """Общая часть имени файлов сессии."""
    return f"{_device_folder(device_name, device_id)}_{session_stamp(moment)}"


def _timestamp() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


class RawLogWriter:
    """Пишет журнал сессии построчно, деля его на части."""

    def __init__(
        self,
        directory: Path,
        base_name: str,
        *,
        max_lines: int = RAW_MAX_LINES,
        max_bytes: int = RAW_MAX_BYTES,
    ):
        self.directory = Path(directory)
        self.directory.mkdir(parents=True, exist_ok=True)
        self.base_name = base_name
        self.max_lines = max_lines
        self.max_bytes = max_bytes
        self.total_lines = 0
        self._part_index = 0
        self._current_path: Path | None = None
        self._current_file = None
        self._current_lines = 0
        self._current_bytes = 0
        self._parts: list[Path] = []
        self._open_new_part()

    def _open_new_part(self) -> None:
        if self._current_file is not None:
            self._current_file.close()
        self._part_index += 1
        self._current_path = (
            self.directory / f"{self.base_name}_part{self._part_index:03d}.txt"
        )
        self._current_file = open(
            self._current_path, "a", encoding="utf-8", newline="\n"
        )
        self._parts.append(self._current_path)
        self._current_lines = 0
        self._current_bytes = self._current_path.stat().st_size

    def write_line(self, text: str, prefix_timestamp: bool = True) -> None:
        if self._current_file is None:
            raise RuntimeError("Журнал сессии уже закрыт")
        line = f"[{_timestamp()}] {text}" if prefix_timestamp else text
        line = line.rstrip("\n") + "\n"
        self._current_file.write(line)
        self._current_file.flush()
        self._current_lines += 1
        self.total_lines += 1
        # Размер считается по записанным байтам: обращаться к stat на каждой
        # строке дорого, а поток может идти сотнями строк в секунду.
        self._current_bytes += len(line.encode("utf-8"))
        if self._current_lines >= self.max_lines or self._current_bytes >= self.max_bytes:
            self._open_new_part()

    def close(self) -> None:
        if self._current_file is not None:
            self._current_file.close()
            self._current_file = None

    @property
    def parts(self) -> list[Path]:
        return list(self._parts)

    @property
    def current_path(self) -> Path | None:
        return self._current_path


def _iter_lines(parts: list[Path]):
    for part in parts:
        with open(part, "r", encoding="utf-8", errors="replace") as source:
            for line in source:
                yield line.rstrip("\n")


def split_timestamp(line: str) -> tuple[str, str]:
    """Отделяет метку времени от текста строки журнала."""
    if line.startswith("[") and "] " in line:
        stamp, rest = line[1:].split("] ", 1)
        return stamp, rest
    return "", line


def export_session(
    parts: list[Path],
    destination: Path,
    export_format: str,
    base_name: str,
    *,
    docx_max_lines: int = DOCX_MAX_LINES,
    xlsx_max_rows: int = XLSX_MAX_ROWS,
) -> list[Path]:
    """Сохраняет накопленный журнал в выбранном формате.

    Возвращает список созданных файлов.
    """
    destination = Path(destination)
    destination.mkdir(parents=True, exist_ok=True)
    chosen = export_format.lower()
    if chosen == "txt":
        return _export_txt(parts, destination, base_name)
    if chosen == "docx":
        return _export_docx(parts, destination, base_name, docx_max_lines)
    if chosen == "xlsx":
        return _export_xlsx(parts, destination, base_name, xlsx_max_rows)
    raise ValueError(f"Неизвестный формат экспорта: {export_format}")


def _export_txt(parts: list[Path], destination: Path, base_name: str) -> list[Path]:
    created: list[Path] = []
    for index, part in enumerate(parts, start=1):
        target = destination / f"{base_name}_part{index:03d}.txt"
        shutil.copyfile(part, target)
        created.append(target)
    return created


def _export_docx(
    parts: list[Path],
    destination: Path,
    base_name: str,
    max_lines: int,
) -> list[Path]:
    from docx import Document

    created: list[Path] = []
    document = Document()
    count = 0
    part_index = 1

    def flush() -> None:
        nonlocal part_index
        target = destination / f"{base_name}_part{part_index:03d}.docx"
        document.save(target)
        created.append(target)
        part_index += 1

    for line in _iter_lines(parts):
        document.add_paragraph(line)
        count += 1
        if count >= max_lines:
            flush()
            document = Document()
            count = 0
    if count > 0 or not created:
        flush()
    return created


def _export_xlsx(
    parts: list[Path],
    destination: Path,
    base_name: str,
    max_rows: int,
) -> list[Path]:
    from openpyxl import Workbook

    created: list[Path] = []
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Журнал"
    sheet.append(["Время", "Сообщение"])
    rows = 0
    part_index = 1

    def flush() -> None:
        nonlocal part_index
        target = destination / f"{base_name}_part{part_index:03d}.xlsx"
        workbook.save(target)
        created.append(target)
        part_index += 1

    for line in _iter_lines(parts):
        stamp, message = split_timestamp(line)
        sheet.append([stamp, message])
        rows += 1
        if rows >= max_rows:
            flush()
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Журнал"
            sheet.append(["Время", "Сообщение"])
            rows = 0
    if rows > 0 or not created:
        flush()
    return created
