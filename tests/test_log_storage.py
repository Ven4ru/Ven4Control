import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from ven4control.log_storage import (
    RawLogWriter,
    export_session,
    safe_name,
    session_base_name,
    session_directory,
    split_timestamp,
)


class NamingTests(unittest.TestCase):
    def test_unsafe_characters_are_replaced(self) -> None:
        self.assertEqual("Router_home", safe_name("Router: home"))
        self.assertEqual("device", safe_name("  "))

    def test_cyrillic_names_are_preserved(self) -> None:
        # Большинство устройств в проекте названы по-русски — без явного
        # диапазона а-яА-ЯёЁ такое имя целиком запрещённое и схлопывалось
        # в generic "device" (найдено живым тестом на двух реальных
        # устройствах: оба кириллических имени писали сессии в одну папку).
        self.assertEqual("Роутер_Ишим", safe_name("Роутер Ишим"))
        self.assertEqual("Ёж", safe_name("Ёж"))

    def test_session_directory_separates_devices_and_runs(self) -> None:
        moment = datetime(2026, 9, 7, 18, 30, 45)
        path = session_directory("Router", moment, Path("C:/logs"))
        self.assertEqual(Path("C:/logs/Router/20260907-183045"), path)
        self.assertEqual("Router_20260907-183045", session_base_name("Router", moment))

    def test_device_id_disambiguates_same_or_colliding_names(self) -> None:
        # Имя — не гарантированно уникальный ключ (пользователь может дать
        # двум устройствам одинаковое имя, или разные имена схлопнутся в
        # один safe_name); id устройства — единственное, что гарантированно
        # разное, поэтому он должен различать папки, если передан.
        moment = datetime(2026, 9, 7, 18, 30, 45)
        first = session_directory("Роутер", moment, Path("C:/logs"), device_id=1)
        second = session_directory("Роутер", moment, Path("C:/logs"), device_id=4)
        self.assertNotEqual(first, second)
        self.assertEqual(Path("C:/logs/Роутер-1/20260907-183045"), first)
        self.assertEqual(Path("C:/logs/Роутер-4/20260907-183045"), second)
        self.assertEqual(
            "Роутер-4_20260907-183045",
            session_base_name("Роутер", moment, device_id=4),
        )


class RawLogWriterTests(unittest.TestCase):
    def test_lines_are_written_with_timestamp(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            writer = RawLogWriter(Path(temporary), "session")
            writer.write_line("строка журнала")
            writer.write_line("=== МАРКЕР ===", prefix_timestamp=False)
            writer.close()

            content = writer.parts[0].read_text(encoding="utf-8").splitlines()
            self.assertRegex(content[0], r"^\[\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}\] строка журнала$")
            self.assertEqual("=== МАРКЕР ===", content[1])
            self.assertEqual(2, writer.total_lines)

    def test_part_is_rotated_by_line_count(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            writer = RawLogWriter(Path(temporary), "session", max_lines=3)
            for number in range(7):
                writer.write_line(f"строка {number}")
            writer.close()

            self.assertEqual(3, len(writer.parts))
            self.assertEqual(
                3, len(writer.parts[0].read_text(encoding="utf-8").splitlines())
            )
            self.assertEqual(
                1, len(writer.parts[2].read_text(encoding="utf-8").splitlines())
            )

    def test_part_is_rotated_by_size(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            writer = RawLogWriter(Path(temporary), "session", max_bytes=120)
            for number in range(6):
                writer.write_line(f"довольно длинная строка номер {number}")
            writer.close()

            self.assertGreater(len(writer.parts), 1)
            for part in writer.parts[:-1]:
                self.assertGreaterEqual(part.stat().st_size, 120)

    def test_parts_are_numbered_from_one(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            writer = RawLogWriter(Path(temporary), "session", max_lines=1)
            writer.write_line("одна")
            writer.write_line("две")
            writer.close()

            self.assertEqual("session_part001.txt", writer.parts[0].name)
            self.assertEqual("session_part002.txt", writer.parts[1].name)

    def test_write_after_close_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            writer = RawLogWriter(Path(temporary), "session")
            writer.close()
            with self.assertRaises(RuntimeError):
                writer.write_line("поздно")

    def test_existing_part_size_is_taken_into_account(self) -> None:
        """Повторное открытие папки сессии не должно терять прежний объём."""
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            (directory / "session_part001.txt").write_text(
                "x" * 200, encoding="utf-8"
            )
            writer = RawLogWriter(directory, "session", max_bytes=120)
            writer.write_line("строка")
            writer.close()

            self.assertEqual(2, len(writer.parts))


class TimestampSplitTests(unittest.TestCase):
    def test_prefixed_line_is_split(self) -> None:
        self.assertEqual(
            ("2026-09-07 18:00:00", "текст"),
            split_timestamp("[2026-09-07 18:00:00] текст"),
        )

    def test_line_without_prefix_stays_whole(self) -> None:
        self.assertEqual(("", "=== СНАПШОТ ==="), split_timestamp("=== СНАПШОТ ==="))


class ExportTests(unittest.TestCase):
    def _session(self, temporary: str, lines: int = 5) -> RawLogWriter:
        writer = RawLogWriter(Path(temporary) / "raw", "session")
        for number in range(lines):
            writer.write_line(f"строка {number}")
        writer.close()
        return writer

    def test_unknown_format_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            with self.assertRaises(ValueError):
                export_session([], Path(temporary), "pdf", "session")

    def test_txt_export_copies_parts(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            writer = self._session(temporary)
            destination = Path(temporary) / "export"

            created = export_session(writer.parts, destination, "txt", "session")

            self.assertEqual(1, len(created))
            self.assertEqual(
                writer.parts[0].read_text(encoding="utf-8"),
                created[0].read_text(encoding="utf-8"),
            )

    def test_docx_export_is_split_by_threshold(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            writer = self._session(temporary, lines=5)
            destination = Path(temporary) / "export"

            created = export_session(
                writer.parts, destination, "docx", "session", docx_max_lines=2
            )

            self.assertEqual(3, len(created))
            self.assertTrue(all(item.exists() for item in created))
            from docx import Document

            first = Document(str(created[0]))
            self.assertEqual(2, len(first.paragraphs))
            self.assertIn("строка 0", first.paragraphs[0].text)

    def test_xlsx_export_splits_timestamp_into_columns(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            writer = self._session(temporary, lines=3)
            destination = Path(temporary) / "export"

            created = export_session(writer.parts, destination, "xlsx", "session")

            self.assertEqual(1, len(created))
            from openpyxl import load_workbook

            sheet = load_workbook(created[0]).active
            self.assertEqual(["Время", "Сообщение"], [sheet["A1"].value, sheet["B1"].value])
            self.assertEqual("строка 0", sheet["B2"].value)
            self.assertRegex(str(sheet["A2"].value), r"^\d{4}-\d{2}-\d{2} ")

    def test_empty_session_still_produces_a_file(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            destination = Path(temporary) / "export"
            for export_format in ("docx", "xlsx"):
                with self.subTest(export_format=export_format):
                    created = export_session([], destination, export_format, "empty")
                    self.assertEqual(1, len(created))
                    self.assertTrue(created[0].exists())


if __name__ == "__main__":
    unittest.main()
